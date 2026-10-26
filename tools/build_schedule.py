"""Rebuild miniprogram/data/events.js from the finalized Excel workbook.
Run from project root: python tools/build_schedule.py

The canonical input sheet is `有课日期速查`.
`Tutorial候选` = precise selectable group candidate (currently PMAA7001/7003).
`Tutorial时段` = timetable-only candidate window without a confirmed group/sub-slot;
                     it stays visible but is nonBlocking in runtime schedule logic.
"""
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime
import json
import re

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'source' / 'sample-timetable-2026-27.xlsx'
OUT = ROOT / 'miniprogram' / 'data' / 'events.js'
COURSES = ROOT / 'miniprogram' / 'data' / 'courses.js'


def load_js_json(path):
    raw = path.read_text(encoding='utf-8').strip()
    if raw.startswith('module.exports = '):
        raw = raw[len('module.exports = '):]
    if raw.endswith(';'):
        raw = raw[:-1]
    return json.loads(raw)


course_map = {c['code']: c for c in load_js_json(COURSES)}
existing = load_js_json(OUT) if OUT.exists() else []


def norm_time(s):
    if not s or not isinstance(s, str):
        return None, None
    s = s.replace('—', '–')
    if '–' not in s:
        return None, None
    a, b = s.split('–', 1)
    return a.strip(), b.strip()


def etype(t, code):
    return {
        '大课': 'lecture',
        '考试': 'exam',
        'Tutorial候选': 'tutorial',
        'Tutorial时段': 'tutorial',
        'Boot Camp': 'bootcamp',
        '待确认': 'tbc',
    }.get(t, {
        'ORIENT': 'orientation',
        'CAREER': 'career',
        'MODULE6': 'module-marker',
    }.get(code, 'calendar'))


def signature(e):
    return (
        e.get('date'), e.get('courseCode'), e.get('type'),
        e.get('startTime'), e.get('endTime'), e.get('tutorialGroup'),
        e.get('title') or '',
    )


# Reuse IDs for unchanged occurrences so local temporary-reschedule records keep working.
existing_ids = {signature(e): e.get('id') for e in existing if e.get('id')}

wb = load_workbook(SRC, data_only=True)
ws = wb['有课日期速查']
out = []
for row in ws.iter_rows(min_row=4, values_only=True):
    d, wd, stage, ts, item, t, venue, note = row
    if not d or not item:
        continue
    ds = d.strftime('%Y-%m-%d') if isinstance(d, datetime) else str(d)
    item_text = str(item)
    if '｜' in item_text:
        code, title = item_text.split('｜', 1)
    else:
        code, title = '', item_text
    code = code.strip()
    title = title.strip()
    start, end = norm_time(ts)
    m = re.search(r'Tutorial\s+([A-Z][0-9]+)', title, re.I)
    group = m.group(1).upper() if m else None
    typ = etype(t, code)
    non_blocking = (t == 'Tutorial时段')
    ev = {
        'date': ds,
        'weekday': wd or '',
        'module': stage or '',
        'type': typ,
        'courseCode': code if code.startswith('PMAA') or code.startswith('BOOT-') else None,
        'code': code,
        'title': title,
        'startTime': start,
        'endTime': end,
        'venue': venue or '',
        'tutorialGroup': group,
        'status': 'tbc' if t == '待确认' else 'confirmed',
        'note': note or '',
        'color': course_map.get(code, {}).get('color', '#94A3B8'),
    }
    if non_blocking:
        ev['nonBlocking'] = True
        ev['keepTitle'] = True
    sig = signature(ev)
    ev['id'] = existing_ids.get(sig) or re.sub(
        r'[^A-Za-z0-9_-]', '',
        f'{ds}-{code}-{start or t}-{group or ""}-{len(out)}'
    )
    # Keep id first for a stable, readable generated JS file.
    ev = {'id': ev['id'], **{k: v for k, v in ev.items() if k != 'id'}}
    out.append(ev)

OUT.write_text('module.exports = ' + json.dumps(out, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
print(f'Wrote {len(out)} events -> {OUT}')
print('Non-blocking tutorial windows:', sum(1 for e in out if e.get('nonBlocking')))
